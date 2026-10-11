#!/usr/bin/env python3
"""Five-table, offline XLSX adapted from rebuilder's workbook utility."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from _contract import load_plan, WORKBOOK, need

TABLES={
 '原稿视觉DNA':['维度','对象/区域','观察与详细判断','证据','置信度','教学影响'],
 '色值与对象':['对象类别','对象','色值职责','HEX','RGB','采样/组合来源','证据/近似性'],
 '十步画法':['编号','PNG文件','学习目标','前置状态','对象/区域','图层关系','工具/笔刷职责',
            '笔触行为','执行动作','可见变化','完成后状态','完成标准','图上标题','图上操作',
            '图上笔触','图上检查','过程证据','设置证据','来源ID','视图'],
 '笔刷与操作':['职责ID','职责','原生名称','目标笔触','参数/建议','证据','依据/验证条件','来源ID','替代方法'],
 '资料来源':['来源ID','类别','链接/路径','提交/版本','内容哈希','实际读取范围','授权','采用方式','限制']}

def table_data(plan):
    dna=[[r[k] for k in ('dimension','region','observation','evidence','confidence','teaching_effect')] for r in plan['visual_dna']]
    dna += [['原稿标识','整体',plan['source']['title']+' | '+plan['source']['sha256'],'OBSERVED','哈希绑定',plan['source']['path']],
            ['生成记录','教程页面',plan['generation']['tool']+' | '+plan['generation']['model'],'OBSERVED',
             plan['generation']['resolution_method'],plan['generation']['limitations']]]
    palette=[]
    for r in plan['palette']:
        c=r['hex'].lstrip('#');rgb=', '.join(str(int(c[i:i+2],16)) for i in (0,2,4))
        palette.append([r['category'],r['object'],r['role'],r['hex'],rgb,r['source'],r['evidence']])
    brush_index={b['id']:b for b in plan['brushes']}
    steps=[]
    for s in plan['steps']:
        t=s['display_text']
        steps.append([s['id'],s['filename'],s['learning_goal'],'；'.join(s['preconditions']),s['region'],s['layer_action'],
                      brush_index[s['brush_id']]['role'],s['brush_behavior'],s['operation'],s['visible_delta'],s['state_after'],
                      s['completion_check'],t['title'],t['action'],t['brush_note'],t['check'],s['process_evidence'],
                      s['setting_evidence'],'；'.join(s['source_basis']),s['view']])
    brushes=[[str(b.get(k) or '未核实原生名称，采用职责说明') if k=='native_name' else b[k]
              for k in ('id','role','native_name','behavior','settings','evidence','evidence_basis','source_id','alternative')]
             for b in plan['brushes']]
    sources=[[r[k] for k in ('id','kind','location','commit','content_hash','read_scope','license','adoption','limit')]
             for r in plan['sources']]
    return {'原稿视觉DNA':dna,'色值与对象':palette,'十步画法':steps,'笔刷与操作':brushes,'资料来源':sources}

def validate_workbook(path,plan):
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        need(wb.sheetnames==list(TABLES),'workbook must have exactly the five specified sheets')
        need(not getattr(wb,'_external_links',[]),'external workbook links forbidden')
        rows=table_data(plan)
        for name,headers in TABLES.items():
            ws=wb[name]
            expected=[headers]+rows[name]
            actual=[[c.value for c in row] for row in ws.iter_rows()]
            need(actual==expected,f'workbook content differs from plan: {name}')
            need(all(c.data_type!='f' for row in ws.iter_rows() for c in row),'formulas forbidden')
        need(wb['十步画法'].max_row==11,'ten steps required in workbook')
    finally:
        wb.close()

def build(plan_path,output):
    plan=load_plan(plan_path)
    template=Path(__file__).resolve().parents[1]/'assets'/'原稿风格与画法资料.template.xlsx'
    wb=load_workbook(template)
    need(wb.sheetnames==list(TABLES),'workbook template schema mismatch')
    rows=table_data(plan)
    for name,headers in TABLES.items():
        ws=wb[name]
        if ws.max_row>1:
            ws.delete_rows(2,ws.max_row-1)
        for i,heading in enumerate(headers,1):
            cell=ws.cell(1,i,heading);cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='244A50')
        for row in rows[name]:
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.data_type='s'  # User strings are text, never executable formulas.
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    output=Path(output)
    need(not output.is_symlink(),'output workbook cannot be a symlink')
    output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output);wb.close()
    validate_workbook(output,plan)
    return {'status':'WORKBOOK_PASS','output':str(output),'sheets':5,'steps':10}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('plan');ap.add_argument('workdir',type=Path)
    a=ap.parse_args();print(json.dumps(build(a.plan,a.workdir/WORKBOOK),ensure_ascii=False))

if __name__=='__main__':
    main()
