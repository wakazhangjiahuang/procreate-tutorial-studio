"""Synthetic fixtures only. These tests do not establish artistic/iPad quality."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from PIL import Image, ImageCms, ImageDraw
from openpyxl import load_workbook

SCRIPTS=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SCRIPTS))
from _contract import ContractError, DNA_DIMENSIONS, PAGE_GATES, TRANSITION_GATES, WORKBOOK, digest, load_plan, validate_plan
from verify_pages import verify
from build_workbook import build, validate_workbook
from package_delivery import package
from select_assets import select

class DeliveryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix='procreate-tutorial-test-')
        cls.root=Path(cls.temp.name);cls.pages=cls.root/'pages';cls.pages.mkdir()
        cls.profile=ImageCms.ImageCmsProfile(ImageCms.createProfile('sRGB')).tobytes()
        for i in range(1,11):
            im=Image.new('RGB',(3072,4096),'white');draw=ImageDraw.Draw(im)
            for j in range(i):
                draw.rectangle((400+(j%5)*400,700+(j//5)*600,650+(j%5)*400,950+(j//5)*600),fill=(40+j*10,100,130))
            im.save(cls.pages/f'{i:02d}_测试形块{i}.png',icc_profile=cls.profile)
        cls.source=cls.root/'master.png';cls.source.write_bytes((cls.pages/'10_测试形块10.png').read_bytes())
        cls.plan={'schema_version':'1.1.0','work_type':'SYNTHETIC_TEST',
                  'source':{'id':'artwork_master','title':'合成几何测试，非用户插画','path':str(cls.source),
                            'sha256':digest(cls.source),'process_evidence':'REASONED'},
                  'delivery':{'count':10,'width':3072,'height':4096,'format':'PNG','color_space':'sRGB','background':'#FFFFFF'},
                  'generation':{'tool':'Pillow geometric fixture','model':'NO_AI_MODEL','resolution_method':'Native synthetic 3072x4096','limitations':'CODE TEST ONLY; no artistic validation'},
                  'visual_dna':[{'dimension':d,'region':'合成图形','observation':'用于验证字段与文件，不归因艺术风格',
                                 'evidence':'OBSERVED','confidence':'代码测试','teaching_effect':'仅验证结构'} for d in sorted(DNA_DIMENSIONS)],
                  'palette':[{'category':'主体','object':'测试矩形','role':'测试色','hex':'#286482','source':'代码常量','evidence':'OBSERVED'}],
                  'brushes':[{'id':'block','role':'测试色块职责','native_name':None,'behavior':'建立矩形',
                              'settings':'测试几何尺寸，不是 Procreate 参数','evidence':'PROPOSED','evidence_basis':'仅合成数据',
                              'source_id':'fixture','alternative':'仅测试'}],
                  'sources':[{'id':'fixture','kind':'SYNTHETIC_TEST','location':'local fixture generator','commit':'N/A',
                              'content_hash':digest(cls.source),'read_scope':'code-generated shapes','license':'self-generated test',
                              'adoption':'test fixture','limit':'no real artwork or drawing trial'}],
                  'operations':[{'id':f'op{i}','classification':'REQUIRED','goal':f'测试形块{i}',
                                 'object_material':'synthetic geometry','depends_on':[f'op{i-1}'] if i>1 else []} for i in range(1,11)],
                  'steps':[]}
        for i in range(1,11):
            cls.plan['steps'].append({'id':f'{i:02d}','type':'step','filename':f'{i:02d}_测试形块{i}.png','action_type':'FILL',
              'learning_goal':f'建立测试形块{i}','preconditions':['白画布与前序已成立'],'region':f'测试位置{i}',
              'layer_action':'测试图形层','brush_id':'block','brush_behavior':'矩形填充','operation':f'建立第{i}个矩形',
              'visible_delta':f'新增形块{i}','state_after':f'合成图形{i}','completion_check':'矩形存在',
              'view':'固定全景','operation_ids':[f'op{i}'],'source_basis':['fixture'],'process_evidence':'REASONED',
              'setting_evidence':'PROPOSED','references':['artwork_master']+([f'{i-1:02d}'] if i>1 else []),
              'white_regions':[[0,0,200,4096]],'display_text':{'title':f'{i:02d} 建立形块','action':'测试填充动作',
                                                         'brush_note':'合成几何，非笔刷试画','check':'检查形块存在'}})
        cls.plan_path=cls.root/'plan.json';cls.write_json(cls.plan_path,cls.plan)
        cls.review={'plan_sha256':digest(cls.plan_path),'reviewer':'SYNTHETIC_CODE_TEST','review_method':'STRUCTURAL_FIXTURE_RECORD',
                    'reviewed_at':'2026-10-11','scope':'SYNTHETIC_TEST_ONLY','first_step_startable':'PASS','terminal_matches_master':'PASS',
                    'pages':[{'id':s['id'],'sha256':digest(cls.pages/s['filename']),'basis':'仅合成代码记录，非艺术结论',
                              **{g:'PASS' for g in PAGE_GATES}} for s in cls.plan['steps']],
                    'transitions':[{'from':f'{i:02d}','to':f'{i+1:02d}','basis':'仅测试序列字段',
                                    **{g:'PASS' for g in TRANSITION_GATES}} for i in range(1,10)]}
        cls.review_path=cls.root/'review.json';cls.write_json(cls.review_path,cls.review)

    @staticmethod
    def write_json(path,data):
        Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_complete_white_4k_to_five_sheets_to_11_member_zip(self):
        load_plan(self.plan_path);verify(self.plan_path,self.pages,self.review_path)
        build(self.plan_path,self.pages/WORKBOOK)
        (self.pages/'QA报告.txt').write_text('must not be packaged')
        result=package(self.plan_path,self.pages,self.review_path,self.root/'delivery.zip')
        self.assertEqual(result['status'],'PACKAGE_TEST_PASS')
        with zipfile.ZipFile(self.root/'delivery.zip') as z:
            self.assertEqual(z.namelist(),[s['filename'] for s in self.plan['steps']]+[WORKBOOK])
            self.assertEqual(len(z.infolist()),11)

    def test_missing_required_operation_and_future_dependency(self):
        for change in ('missing','future'):
            p=copy.deepcopy(self.plan)
            if change=='missing': p['operations'].append({'id':'missing','classification':'REQUIRED','goal':'缺口','object_material':'test','depends_on':[]})
            else: p['operations'][0]['depends_on']=['op10']
            with self.assertRaises(ContractError): validate_plan(p)

    def test_cover_summary_and_no_action_rejected(self):
        for change in ('cover','summary','inspect'):
            p=copy.deepcopy(self.plan)
            if change=='cover': p['steps'][0]['type']='cover'
            elif change=='summary': p['steps'][9]['display_text']['title']='成品展示'
            else: p['steps'][9]['action_type']='INSPECT'
            with self.assertRaises(ContractError): validate_plan(p)

    def test_unsafe_filename_missing_master_and_duplicate_operation(self):
        for change in ('escape','master','operation'):
            p=copy.deepcopy(self.plan)
            if change=='escape': p['steps'][0]['filename']='../escape.png'
            elif change=='master': p['steps'][0]['references']=[]
            else: p['steps'][1]['operation_ids']=['op1']
            with self.assertRaises(ContractError): validate_plan(p)

    def test_native_name_without_native_evidence_rejected(self):
        p=copy.deepcopy(self.plan);p['brushes'][0]['native_name']='Invented Native Brush'
        with self.assertRaises(ContractError): validate_plan(p)

    def test_review_not_run_stale_plan_missing_transition_rejected(self):
        for change in ('not_run','stale','missing'):
            r=copy.deepcopy(self.review)
            if change=='not_run': r['pages'][0]['text_exact']='NOT_RUN'
            elif change=='stale': r['plan_sha256']='0'*64
            else: r['transitions'].pop()
            path=self.root/f'review-{change}.json';self.write_json(path,r)
            with self.assertRaises(ContractError): verify(self.plan_path,self.pages,path)

    def test_dimensions_color_white_and_hash_mutation(self):
        target=self.pages/self.plan['steps'][0]['filename'];original=target.read_bytes()
        try:
            Image.new('RGB',(30,40),'white').save(target,icc_profile=self.profile)
            with self.assertRaises(ContractError): verify(self.plan_path,self.pages)
            Image.new('RGB',(3072,4096),'white').save(target)
            with self.assertRaises(ContractError): verify(self.plan_path,self.pages)
            Image.new('RGB',(3072,4096),(250,250,250)).save(target,icc_profile=self.profile)
            with self.assertRaises(ContractError): verify(self.plan_path,self.pages)
            Image.new('RGB',(3072,4096),'white').save(target,icc_profile=self.profile)
            with self.assertRaises(ContractError): verify(self.plan_path,self.pages,self.review_path)
        finally:
            target.write_bytes(original)

    def test_workbook_extra_sheet_wrong_reference_and_formula_rejected(self):
        path=self.root/'mutated.xlsx'
        for change in ('extra','wrong_reference','formula'):
            build(self.plan_path,path)
            wb=load_workbook(path)
            if change=='extra': wb.create_sheet('QA')
            elif change=='wrong_reference': wb['十步画法']['B2']='wrong.png'
            else: wb['资料来源']['A2']='=1+1'
            wb.save(path);wb.close()
            with self.assertRaises(ContractError): validate_workbook(path,self.plan)

    def test_source_changed_rejected(self):
        p=copy.deepcopy(self.plan);p['source']['sha256']='0'*64
        with self.assertRaises(ContractError): validate_plan(p)

    def test_manifest_select_and_integrity_failure(self):
        source=self.root/'asset.md';source.write_text('audited asset')
        manifest=self.root/'KB-MANIFEST.json'
        self.write_json(manifest,{'schema_version':'1.1.0','entrypoints':{'tutorial':['rule']},
                                'assets':[{'id':'rule','path':'asset.md','sha256':digest(source),'task_tags':['tutorial']}]})
        self.assertEqual(len(select(manifest,'tutorial')),1)
        source.write_text('changed')
        with self.assertRaises(ContractError): select(manifest,'tutorial')

if __name__=='__main__':
    unittest.main(verbosity=2)
